#!/usr/bin/env python3
"""Render docs/publishing-deck.md as a .pptx, script into the speaker notes.

The markdown is the source of truth. Each `## N — Title · X min` becomes a
slide; everything up to `**Script**` becomes the slide body; everything after
becomes the notes.
"""

import re
import sys
import pathlib

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

SRC = pathlib.Path(sys.argv[1])
OUT = pathlib.Path(sys.argv[2])

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
GREEN = RGBColor(0x76, 0xB9, 0x00)          # NVIDIA green
CODE_BG = RGBColor(0xF4, 0xF5, 0xF6)
RULE = RGBColor(0xD8, 0xDA, 0xDC)

BODY_FONT = "Segoe UI"
MONO_FONT = "Consolas"

# ---------------------------------------------------------------- parsing ---


def parse(text):
    slides = []
    chunks = re.split(r"\n(?=## )", text)
    for chunk in chunks:
        m = re.match(r"## ([0-9]+|B[0-9]+) — ([^\n·]+?)(?:·\s*([^\n]+))?\n", chunk)
        if not m:
            continue
        num, title, timing = m.group(1), m.group(2).strip(), (m.group(3) or "").strip()
        rest = chunk[m.end():]
        if "**Script**" in rest:
            body, notes = rest.split("**Script**", 1)
        else:
            body, notes = rest, ""
        notes = re.sub(r"\n---\s*$", "", notes.strip())
        notes = re.sub(r"\n#.*$", "", notes, flags=re.S).strip()
        slides.append(dict(num=num, title=title, timing=timing,
                           blocks=blocks_of(body.strip()), notes=clean_notes(notes)))
    return slides


def clean_notes(s):
    s = re.sub(r"`([^`]+)`", r"\1", s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
    s = re.sub(r"\*([^*]+)\*", r"\1", s)
    return s.strip()


def blocks_of(body):
    """-> list of ('code'|'table'|'head'|'bullet'|'para', payload)"""
    out, i = [], 0
    lines = body.split("\n")
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            j = i + 1
            buf = []
            while j < len(lines) and not lines[j].startswith("```"):
                buf.append(lines[j]); j += 1
            out.append(("code", buf))
            i = j + 1
        elif line.startswith("|"):
            buf = []
            while i < len(lines) and lines[i].startswith("|"):
                buf.append(lines[i]); i += 1
            out.append(("table", buf))
        elif line.startswith("### "):
            out.append(("head", inline(line[4:])))
            i += 1
        elif line.startswith("- "):
            buf = []
            while i < len(lines) and lines[i].startswith("- "):
                buf.append(inline(lines[i][2:])); i += 1
            out.append(("bullet", buf))
        elif line.strip():
            buf = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not lines[i].startswith(("```", "|", "- ", "### ")):
                buf.append(lines[i]); i += 1
            out.append(("para", inline(" ".join(buf))))
        else:
            i += 1
    return out


def inline(s):
    """Return [(text, bold, mono)] runs."""
    runs, pos = [], 0
    for m in re.finditer(r"\*\*(.+?)\*\*|`(.+?)`|\*(.+?)\*", s):
        if m.start() > pos:
            runs.append((s[pos:m.start()], False, False))
        if m.group(1) is not None:
            runs.append((m.group(1), True, False))
        elif m.group(2) is not None:
            runs.append((m.group(2), False, True))
        else:
            runs.append((m.group(3), False, False))
        pos = m.end()
    if pos < len(s):
        runs.append((s[pos:], False, False))
    return runs or [(s, False, False)]


def table_rows(lines):
    rows = []
    for ln in lines:
        if re.fullmatch(r"\|[\s:|]*-[\s:\-|]*\|", ln.strip()):
            continue
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


# --------------------------------------------------------------- rendering --

W, H = Inches(13.333), Inches(7.5)
L, R = Inches(0.62), Inches(0.62)
CONTENT_W = W - L - R


def add_runs(p, runs, size, color=INK):
    p.alignment = PP_ALIGN.LEFT
    for text, bold, mono in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = MONO_FONT if mono else BODY_FONT
        r.font.color.rgb = color


def build(slides, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # ---- title slide ----
    s = prs.slides.add_slide(blank)
    box = s.shapes.add_textbox(L, Inches(2.5), CONTENT_W, Inches(2.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    add_runs(p, [("Splitting the CCCL docs", True, False)], 40)
    p2 = tf.add_paragraph()
    add_runs(p2, [("Two products, two release cadences, two documentation sites", False, False)], 18, MUTED)
    p2.space_before = Pt(14)
    p3 = tf.add_paragraph()
    add_runs(p3, [("https://sri-koundinyan.github.io/cccl/", False, True)], 14, GREEN)
    p3.space_before = Pt(18)

    for sl in slides:
        s = prs.slides.add_slide(blank)

        # title
        tb = s.shapes.add_textbox(L, Inches(0.42), CONTENT_W, Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        add_runs(p, [(sl["title"], True, False)], 26)
        if sl["timing"]:
            tb2 = s.shapes.add_textbox(W - R - Inches(1.4), Inches(0.52), Inches(1.4), Inches(0.35))
            q = tb2.text_frame.paragraphs[0]
            q.alignment = PP_ALIGN.RIGHT
            add_runs(q, [(sl["timing"], False, False)], 11, MUTED)

        ln = s.shapes.add_shape(1, L, Inches(1.16), CONTENT_W, Emu(9525))
        ln.fill.solid(); ln.fill.fore_color.rgb = RULE
        ln.line.fill.background(); ln.shadow.inherit = False

        y = Inches(1.42)
        for kind, payload in sl["blocks"]:
            y = render_block(s, kind, payload, y)

        # notes
        if sl["notes"]:
            s.notes_slide.notes_text_frame.text = sl["notes"]

    prs.save(out)
    return len(prs.slides._sldIdLst)


def render_block(s, kind, payload, y):
    if kind == "code":
        lines = payload
        size = 12 if max((len(l) for l in lines), default=0) <= 84 else 10
        h = Inches(0.24) * len(lines) + Inches(0.26)
        bg = s.shapes.add_shape(1, L, y, CONTENT_W, h)
        bg.fill.solid(); bg.fill.fore_color.rgb = CODE_BG
        bg.line.fill.background(); bg.shadow.inherit = False
        tb = s.shapes.add_textbox(L + Inches(0.16), y + Inches(0.1), CONTENT_W - Inches(0.3), h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = line or " "
            r.font.size = Pt(size); r.font.name = MONO_FONT; r.font.color.rgb = INK
            p.space_after = Pt(0)
        return y + h + Inches(0.18)

    if kind == "table":
        rows = table_rows(payload)
        headerless = not any(c.strip() for c in rows[0])
        if headerless:
            rows = rows[1:]
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        h = Inches(0.34) * len(rows)
        shp = s.shapes.add_table(len(rows), ncol, L, y, CONTENT_W, h)
        tbl = shp.table
        tbl.first_row = not headerless
        tbl.horz_banding = False
        for ri, row in enumerate(rows):
            tbl.rows[ri].height = Inches(0.34)
            is_head = (ri == 0 and not headerless)
            for ci, cell in enumerate(row):
                c = tbl.cell(ri, ci)
                c.margin_left = Inches(0.1); c.margin_right = Inches(0.08)
                c.margin_top = Inches(0.03); c.margin_bottom = Inches(0.03)
                c.fill.solid()
                c.fill.fore_color.rgb = CODE_BG if is_head else RGBColor(0xFF, 0xFF, 0xFF)
                p = c.text_frame.paragraphs[0]
                add_runs(p, inline(cell), 13)
                if is_head:
                    for r_ in p.runs:
                        r_.font.bold = True
        return y + h + Inches(0.22)

    if kind == "head":
        tb = s.shapes.add_textbox(L, y, CONTENT_W, Inches(0.42))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]
        add_runs(p, payload, 18, GREEN)
        for r in p.runs:
            r.font.bold = True
        return y + Inches(0.52)

    if kind == "bullet":
        h = Inches(0.32) * len(payload) + Inches(0.08)
        tb = s.shapes.add_textbox(L, y, CONTENT_W, h)
        tf = tb.text_frame; tf.word_wrap = True
        for i, runs in enumerate(payload):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r0 = p.add_run(); r0.text = "— "
            r0.font.size = Pt(16); r0.font.name = BODY_FONT; r0.font.color.rgb = MUTED
            add_runs(p, runs, 16)
            p.space_after = Pt(4)
        return y + h + Inches(0.12)

    # para
    tb = s.shapes.add_textbox(L, y, CONTENT_W, Inches(0.5))
    tf = tb.text_frame; tf.word_wrap = True
    add_runs(tf.paragraphs[0], payload, 16)
    nchars = sum(len(t) for t, _, _ in payload)
    return y + Inches(0.34) * max(1, (nchars // 110) + 1) + Inches(0.12)


slides = parse(SRC.read_text(encoding="utf-8"))
n = build(slides, OUT)
print(f"  {len(slides)} content slides + 1 title = {n} slides")
for sl in slides:
    print(f"    {sl['num']:>3}  {sl['title'][:52]:<54} notes {len(sl['notes']):>5} chars")
